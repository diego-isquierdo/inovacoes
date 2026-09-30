"""Extração ANUAL das horas apontadas no Painel de Serviços v2 (Projuris).

Gera UMA planilha do ano com todos os campos que a API devolve, para o racional
de horas e a S&OP do Painel de Gestão de Serviços Técnicos. Planejamento, regras e
histórico de decisões em:
    Painel de Gestão/Serviços Técnicos/extracao/PLANEJAMENTO_EXTRACAO_HORAS.md

Fontes (API REST /api/v2/apontamentos/*), todas somente leitura:
    Apontamentos     GET /apontamentos?mine=false&date_from=AAAA-01-01&date_to=AAAA-12-31   (paginado)
    Áreas            GET /areas?include_inactive=true   (fallback: /me/areas)
    Analistas        GET /analysts?area_id=...
    Membros          GET /members?area_id=...            (liga analista <-> agente Freshdesk)
    Capacidade       GET /reports/capacity?area_id=...&month=AAAA-MM   (12 meses x áreas)
    Fechamento       GET /reports/closing?month=AAAA-MM&area_id=...    (meses até o atual)
    Meta produtiv.   GET /reports/productivity-goal?year_month=AAAA-MM
    Disponibilidade  GET /calendar/availability?year_month=AAAA-07&span=6  (ano inteiro numa chamada)
    Calendário       GET /calendar?year_month=AAAA-MM&analyst_id=...       (dias: esperado x confirmado)

Rascunhos NÃO entram (decisão do Diego, 29/09/2026).

Uma rota que o token não pode ler (401/403/404) não derruba a extração: vira uma
linha na aba "Execução" com o status, e as demais abas saem normalmente.

Configuração (.env na pasta "MCP Painel de Serviços", nunca versionado):
    Reaproveita o .env do MCP (mcp_server/.env):
    PAINEL_SERVICOS_BASE_URL=https://servicos.projuris.com.br   (sem /api/v2)
    PAINEL_SERVICOS_TOKEN=<token Bearer>

Uso:
    python scripts/extracao_horas_apontadas.py              # ano corrente
    python scripts/extracao_horas_apontadas.py --ano 2026
"""

from __future__ import annotations

import argparse
import json
import os
import sys
import time
from datetime import date, datetime, timezone
from pathlib import Path

import httpx
from dotenv import load_dotenv
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent            # .../MCP Painel de Serviços
OUTPUT_DIR = ROOT / "output"
LOGS_DIR = ROOT / "scripts" / "logs"
REGISTRY = ROOT / "scripts" / "registry.jsonl"
PAGE_SIZE = 100
TIMEOUT = 60.0

try:
    sys.stdout.reconfigure(encoding="utf-8")  # console do Windows (cp1252)
except (AttributeError, ValueError):
    pass


# ---------------------------------------------------------------------------
# Log simples (arquivo + console) e registro de execuções (só metadados)
# ---------------------------------------------------------------------------
class Log:
    def __init__(self) -> None:
        LOGS_DIR.mkdir(parents=True, exist_ok=True)
        self.path = LOGS_DIR / f"extracao_horas_apontadas_{datetime.now():%Y%m%d_%H%M%S}.log"
        self.fh = self.path.open("w", encoding="utf-8")

    def __call__(self, msg: str, *args) -> None:
        line = f"{datetime.now():%Y-%m-%d %H:%M:%S} {msg % args if args else msg}"
        print(line)
        self.fh.write(line + "\n")
        self.fh.flush()


# ---------------------------------------------------------------------------
# Cliente HTTP — só GET; retry limitado em 429/5xx/rede
# ---------------------------------------------------------------------------
class Api:
    def __init__(self, base_url: str, token: str, log: Log) -> None:
        self.c = httpx.Client(base_url=base_url.rstrip("/") + "/api/v2",
                              headers={"Authorization": f"Bearer {token}", "Accept": "application/json"},
                              timeout=TIMEOUT)
        self.log = log
        self.calls = 0
        self.status: list[dict] = []   # uma linha por rota/consulta, vai para a aba Execução

    def get(self, path: str, params: dict | None = None, label: str | None = None):
        params = {k: v for k, v in (params or {}).items() if v is not None}
        path = "/apontamentos" + path  # todas as rotas usadas vivem sob /api/v2/apontamentos
        for attempt in range(5):
            try:
                r = self.c.get(path, params=params)
            except httpx.HTTPError as exc:
                self.log("rede em %s (%s) — tentativa %d/5", path, exc.__class__.__name__, attempt + 1)
                time.sleep(5 * (attempt + 1))
                continue
            self.calls += 1
            if r.status_code == 429 or r.status_code >= 500:
                wait = int(r.headers.get("retry-after", "0") or 0) or 10 * (attempt + 1)
                self.log("HTTP %s em %s — aguardando %ss", r.status_code, path, wait)
                time.sleep(wait)
                continue
            ok = r.status_code == 200
            if label:
                self.status.append({"consulta": label, "rota": path, "parametros": json.dumps(params, ensure_ascii=False),
                                    "http": r.status_code, "ok": ok,
                                    "detalhe": "" if ok else r.text[:300]})
            if not ok:
                if r.status_code in (401, 403, 404, 422):
                    self.log("HTTP %s em %s %s — seguindo sem esta consulta", r.status_code, path, params)
                    return None
                r.raise_for_status()
            return r.json()
        raise RuntimeError(f"Desisti de {path} após 5 tentativas")


# ---------------------------------------------------------------------------
# Achatamento genérico: mantém TODOS os campos que a API devolver
# ---------------------------------------------------------------------------
def flat(obj: dict, prefix: str = "") -> dict:
    out = {}
    for k, v in obj.items():
        key = f"{prefix}{k}"
        if isinstance(v, dict):
            out.update(flat(v, key + "."))
        elif isinstance(v, list):
            out[key] = json.dumps(v, ensure_ascii=False) if v else ""
        else:
            out[key] = v
    return out


def months_of(year: int, until_today: bool) -> list[str]:
    last = date.today().month if (until_today and year == date.today().year) else 12
    return [f"{year}-{m:02d}" for m in range(1, last + 1)]


# ---------------------------------------------------------------------------
# Coletas
# ---------------------------------------------------------------------------
def collect_apontamentos(api: Api, year: int, log: Log) -> list[dict]:
    rows, page, total = [], 1, None
    base = {"mine": "false", "date_from": f"{year}-01-01", "date_to": f"{year}-12-31",
            "sort": "date_asc", "page_size": PAGE_SIZE}
    while True:
        data = api.get("", {**base, "page": page}, label="apontamentos" if page == 1 else None)
        if data is None:
            break
        items = data.get("items") or []
        total = data.get("total", total)
        rows.extend(items)
        if page == 1 or page % 10 == 0:
            log("apontamentos: página %d — %d de %s", page, len(rows), total)
        if not items or (total is not None and len(rows) >= total):
            break
        page += 1
    ids = {r.get("id") for r in rows}
    if len(ids) != len(rows):
        log("AVISO: %d apontamentos duplicados na paginação (removidos)", len(rows) - len(ids))
        rows = list({r["id"]: r for r in rows}.values())
    if total is not None and len(rows) != total:
        log("AVISO: API informou total=%s, extraídos %d", total, len(rows))
    return rows


def collect_all(api: Api, year: int, log: Log) -> dict:
    t: dict[str, list[dict]] = {}

    # Áreas
    areas = api.get("/areas", {"include_inactive": "true"}, label="areas")
    if areas is None:
        mine = api.get("/me/areas", label="me/areas") or []
        areas = [{"id": a["area_id"], "key": a.get("area_key"), "name": a.get("area_name"), **a} for a in mine]
    t["Áreas"] = [flat(a) for a in areas]
    area_ids = [a["id"] for a in areas if a.get("id") is not None]
    log("áreas: %d", len(area_ids))

    # Apontamentos (aba principal)
    t["Apontamentos"] = collect_apontamentos(api, year, log)
    log("apontamentos do ano: %d", len(t["Apontamentos"]))

    # Analistas e membros por área
    analysts: dict[str, dict] = {}
    t["Analistas"], t["Membros"] = [], []
    for aid in area_ids:
        for a in api.get("/analysts", {"area_id": aid}, label=f"analysts area={aid}") or []:
            t["Analistas"].append({"area_id": aid, **flat(a)})
            analysts[a["id"]] = a
        for m in api.get("/members", {"area_id": aid}, label=f"members area={aid}") or []:
            t["Membros"].append(flat(m))
    for r in t["Apontamentos"]:  # analistas que só aparecem nos apontamentos
        if r.get("analyst_id") and r["analyst_id"] not in analysts:
            analysts[r["analyst_id"]] = {"id": r["analyst_id"], "name": r.get("analyst_name")}

    # Capacidade (12 meses x área) e meta de produtividade (12 meses)
    t["Capacidade"], t["Meta de produtividade"] = [], []
    for ym in months_of(year, until_today=False):
        for aid in area_ids:
            cap = api.get("/reports/capacity", {"area_id": aid, "month": ym}, label=f"capacity {ym} area={aid}")
            for it in (cap or {}).get("items") or []:
                t["Capacidade"].append({"area_id": aid, "year_month": ym, **flat(it)})
        goal = api.get("/reports/productivity-goal", {"year_month": ym}, label=f"productivity-goal {ym}")
        if goal:
            t["Meta de produtividade"].append(flat(goal))

    # Fechamento mensal por área (meses até o atual)
    for k in ("Fechamento - Totais", "Fechamento - Analistas", "Fechamento - Tipos"):
        t[k] = []
    for ym in months_of(year, until_today=True):
        for aid in area_ids:
            c = api.get("/reports/closing", {"month": ym, "area_id": aid}, label=f"closing {ym} area={aid}")
            if not c:
                continue
            base = {"month": ym, "area_id": aid}
            t["Fechamento - Totais"].append({**base,
                                             **flat(c.get("totals") or {}, "totals."),
                                             **flat(c.get("retrabalho") or {}, "retrabalho."),
                                             **flat(c.get("orcado_realizado") or {}, "orcado_realizado."),
                                             "meta_produtividade_pct": c.get("meta_produtividade_pct")})
            t["Fechamento - Analistas"] += [{**base, **flat(r)} for r in c.get("analysts") or []]
            t["Fechamento - Tipos"] += [{**base, **flat(r)} for r in c.get("tipo_breakdown") or []]

    # Disponibilidade do ano (julho ± 6 meses cobre jan–dez numa chamada)
    t["Alocações"], t["Férias"], t["Disponibilidade"] = [], [], []
    av = api.get("/calendar/availability", {"year_month": f"{year}-07", "span": 6}, label="availability ano")
    for a in (av or {}).get("analysts") or []:
        base = {"analyst_id": a.get("analyst_id"), "analyst_name": a.get("analyst_name"), "area_name": a.get("area_name")}
        t["Disponibilidade"].append({**base, "next_free_date": a.get("next_free_date"),
                                     "range_start": av.get("range_start"), "range_end": av.get("range_end"),
                                     "qtd_alocacoes": len(a.get("allocations") or []),
                                     "qtd_dias_ferias": len(a.get("vacation_days") or [])})
        t["Alocações"] += [{**base, **flat(x)} for x in a.get("allocations") or []]
        t["Férias"] += [{**base, "date": d} for d in a.get("vacation_days") or []]

    # Calendário diário por analista (esperado x confirmado) — meses até o atual
    t["Calendário"] = []
    for an_id, an in analysts.items():
        for ym in months_of(year, until_today=True):
            cal = api.get("/calendar", {"year_month": ym, "analyst_id": an_id}, label=f"calendar {ym} {an_id}")
            for d in (cal or {}).get("days") or []:
                t["Calendário"].append({
                    "analyst_id": an_id, "analyst_name": an.get("name"), "month": ym,
                    "date": d.get("date"), "is_vacation": d.get("is_vacation"),
                    "expected_seconds": d.get("expected_seconds"),
                    "total_confirmed_seconds": d.get("total_confirmed_seconds"),
                    "qtd_apontamentos": len(d.get("apontamentos") or []),
                    "qtd_alocacoes": len(d.get("allocations") or []),
                    **{k: v for k, v in d.items() if k not in
                       ("date", "is_vacation", "expected_seconds", "total_confirmed_seconds",
                        "apontamentos", "allocations", "drafts")},
                })
    return t


# ---------------------------------------------------------------------------
# Saídas
# ---------------------------------------------------------------------------
APONT_FIRST = ["id", "work_date", "analyst_name", "analyst_id", "area_id", "area_name", "status",
               "atividade", "tipo", "produto", "company_id", "company_name",
               "duration_seconds", "duration_hms", "horas",
               "ticket_id", "ticket_url", "ticket_subject", "ticket_status", "ticket_company",
               "freshdesk_conversation_id", "tramite_at", "tramite_body", "tramite_body_html",
               "observacoes", "created_at", "updated_at", "can_edit", "mes"]


PAINEL_DROP = {"tramite_body", "tramite_body_html", "ticket_url", "can_edit"}


def columns_for(rows: list[dict], first: list[str] | None = None) -> list[str]:
    cols = list(first or [])
    for r in rows:
        for k in r:
            if k not in cols:
                cols.append(k)
    return [c for c in cols if any(c in r for r in rows)] if rows else cols


def write_outputs(year: int, tables: dict, summary: dict, status: list[dict]) -> tuple[Path, Path]:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    ap = []
    for r in tables["Apontamentos"]:
        row = flat(r)
        row["horas"] = round((r.get("duration_seconds") or 0) / 3600, 4)  # derivada; segundos originais preservados
        row["mes"] = (r.get("work_date") or "")[:7]
        ap.append(row)
    tables = {"Apontamentos": ap, **{k: v for k, v in tables.items() if k != "Apontamentos"}}

    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in tables.items():
        ws = wb.create_sheet(name[:31])
        cols = columns_for(rows, APONT_FIRST if name == "Apontamentos" else None)
        ws.append(cols or ["(sem dados)"])
        for c in ws[1]:
            c.font = Font(bold=True)
        for r in rows:
            ws.append([r.get(c) for c in cols])
        ws.freeze_panes = "A2"
        if cols:
            ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{max(len(rows) + 1, 2)}"
            for i, c in enumerate(cols, 1):
                w = max([len(str(c))] + [len(str(r.get(c) or "")) for r in rows[:300]]) + 2
                ws.column_dimensions[get_column_letter(i)].width = min(w, 50)
    ws = wb.create_sheet("Execução")
    ws.append(["Item", "Valor"])
    for k, v in summary.items():
        ws.append([k, json.dumps(v, ensure_ascii=False) if isinstance(v, (dict, list)) else v])
    ws.append([])
    ws.append(["consulta", "rota", "parametros", "http", "ok", "detalhe"])
    for s in status:
        ws.append([s["consulta"], s["rota"], s["parametros"], s["http"], s["ok"], s["detalhe"]])
    for row in (ws[1], ws[len(summary) + 3]):
        for c in row:
            c.font = Font(bold=True)
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 40
    ws.column_dimensions["C"].width = 60

    xlsx = OUTPUT_DIR / f"Horas Apontadas {year}.xlsx"
    wb.save(xlsx)

    js = OUTPUT_DIR / f"horas_apontadas_{year}.json"
    payload = {"meta": summary, "tabelas": {}}
    for name, rows in tables.items():
        cols = columns_for(rows, APONT_FIRST if name == "Apontamentos" else None)
        payload["tabelas"][name] = {"colunas": cols, "linhas": [[r.get(c) for c in cols] for r in rows]}
    js.write_text(json.dumps(payload, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")

    # Versão compacta para a carga do painel: sem o texto integral dos trâmites
    # (tramite_body_html/tramite_body respondem por ~75% do tamanho e não entram
    # nas análises). A planilha e o JSON completo continuam com todos os campos.
    slim = {"meta": {**summary, "versao": "painel (sem texto dos trâmites)"}, "tabelas": {}}
    for name, tb in payload["tabelas"].items():
        if name == "Apontamentos":
            keep = [i for i, c in enumerate(tb["colunas"]) if c not in PAINEL_DROP]
            slim["tabelas"][name] = {"colunas": [tb["colunas"][i] for i in keep],
                                     "linhas": [[r[i] for i in keep] for r in tb["linhas"]]}
        else:
            slim["tabelas"][name] = tb
    js_slim = OUTPUT_DIR / f"horas_apontadas_{year}_painel.json"
    js_slim.write_text(json.dumps(slim, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    return xlsx, js


# ---------------------------------------------------------------------------
def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--ano", type=int, default=date.today().year)
    args = ap.parse_args()

    # mesmo .env do MCP painel-servicos (mcp_server/.env) ou um .env próprio na raiz da pasta
    load_dotenv(ROOT / ".env")
    load_dotenv(ROOT / "mcp_server" / ".env")
    url = os.environ.get("PAINEL_SERVICOS_BASE_URL") or os.environ.get("PAINEL_SERVICOS_URL")
    token = os.environ.get("PAINEL_SERVICOS_TOKEN")
    if not url or not token:
        raise SystemExit("Defina PAINEL_SERVICOS_BASE_URL e PAINEL_SERVICOS_TOKEN no .env (mcp_server/.env ou raiz da pasta).")

    log = Log()
    started, t0 = datetime.now(timezone.utc), time.perf_counter()
    status, err, summary = "ok", None, {}
    api = Api(url, token, log)
    try:
        log("extração de horas — ano %d", args.ano)
        tables = collect_all(api, args.ano, log)
        ap_rows = tables["Apontamentos"]
        secs = sum(r.get("duration_seconds") or 0 for r in ap_rows)
        summary = {
            "base": f"Horas Apontadas {args.ano}",
            "gerado_em": datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ"),
            "ano": args.ano,
            "apontamentos": len(ap_rows),
            "horas_total": round(secs / 3600, 2),
            "analistas": len({r.get("analyst_id") for r in ap_rows if r.get("analyst_id")}),
            "areas": sorted({r.get("area_name") for r in ap_rows if r.get("area_name")}),
            "linhas_por_aba": {k: len(v) for k, v in tables.items()},
            "chamadas_api": api.calls,
            "consultas_sem_permissao": sum(1 for s in api.status if not s["ok"]),
        }
        xlsx, js = write_outputs(args.ano, tables, summary, api.status)
        log("salvo: %s e %s", xlsx.name, js.name)
        log("resumo: %s", json.dumps({k: v for k, v in summary.items() if k != "areas"}, ensure_ascii=False))
    except Exception as exc:
        status, err = "error", str(exc)
        log("ERRO: %s", exc)
        raise
    finally:
        dur = round(time.perf_counter() - t0, 2)
        log("concluído em %.2fs (status=%s)", dur, status)
        REGISTRY.parent.mkdir(parents=True, exist_ok=True)
        with REGISTRY.open("a", encoding="utf-8") as fh:
            fh.write(json.dumps({"script": "extracao_horas_apontadas", "started_at": started.isoformat(),
                                 "duration_s": dur, "status": status, "error": err,
                                 "result": {k: v for k, v in summary.items() if k != "areas"}},
                                ensure_ascii=False) + "\n")


if __name__ == "__main__":
    main()
